"""The fixture catalogue — messy spreadsheets + their expected clean output.

The spec calls real, anonymised spreadsheets "the most valuable asset in the
repo". We keep that asset *reproducible*: each fixture is built deterministically
from this module (no committed binaries to drift), and its gabarito — the exact
expected ``data`` and column dtypes — lives right next to it.

Every fixture targets a feature (F1-F10). ``tests/test_end_to_end.py`` runs each
one through ``clean`` and asserts the output equals the gabarito.
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from messy_table import Config


@dataclass(frozen=True)
class Fixture:
    name: str
    feature: str
    build: Callable[[Path], Path]  # writes the file, returns the path
    expected_data: list[dict[str, Any]]
    expected_columns: list[tuple[str, str]]  # (name, dtype)
    config: Config | None = None
    expect_warning: bool = False
    notes: str = ""
    extra: dict[str, Any] = field(default_factory=dict)


# --------------------------------------------------------------------------- #
# Builders                                                                    #
# --------------------------------------------------------------------------- #
def _write_csv(path: Path, text: str, *, encoding: str = "utf-8") -> Path:
    path.write_bytes(text.encode(encoding))
    return path


def _write_xlsx(
    path: Path,
    rows: list[list[Any]],
    *,
    merges: tuple[str, ...] = (),
    bold: tuple[tuple[int, int], ...] = (),
    number_formats: tuple[tuple[int, int, str], ...] = (),
    sheet_title: str = "Plan1",
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    for r, c, fmt in number_formats:
        ws.cell(row=r, column=c).number_format = fmt
    for r, c in bold:
        ws.cell(row=r, column=c).font = Font(bold=True)
    for rng in merges:
        ws.merge_cells(rng)
    wb.save(path)
    return path


# A canonical Excel serial anchor: 45123 == 2023-07-16 (1900 epoch).
_D = dt.date


# --------------------------------------------------------------------------- #
# Catalogue                                                                   #
# --------------------------------------------------------------------------- #
FIXTURES: list[Fixture] = [
    Fixture(
        name="f01_clean_baseline.csv",
        feature="F6/F10",
        build=lambda p: _write_csv(p, "name,age,active\nAlice,30,true\nBob,25,false\n"),
        expected_data=[
            {"name": "Alice", "age": 30, "active": True},
            {"name": "Bob", "age": 25, "active": False},
        ],
        expected_columns=[("name", "str"), ("age", "int"), ("active", "bool")],
        notes="Clean file: no structural fixes, just typing + bool coercion.",
    ),
    Fixture(
        name="f02_title_and_blank.xlsx",
        feature="F1",
        build=lambda p: _write_xlsx(
            p,
            [
                ["Relatório de Vendas 2024"],
                [None],
                ["produto", "qtd"],
                ["Café", 10],
                ["Chá", 5],
            ],
            bold=((3, 1), (3, 2)),
        ),
        expected_data=[{"produto": "Café", "qtd": 10}, {"produto": "Chá", "qtd": 5}],
        expected_columns=[("produto", "str"), ("qtd", "int")],
        notes="Title row + blank row before the header are skipped.",
    ),
    Fixture(
        name="f03_metadata_block.csv",
        feature="F1",
        build=lambda p: _write_csv(
            p,
            "Empresa:,ACME\nData:,2024-01-01\n\nid,nome,valor\n1,Ana,100\n2,Beto,200\n",
        ),
        expected_data=[
            {"id": 1, "nome": "Ana", "valor": 100},
            {"id": 2, "nome": "Beto", "valor": 200},
        ],
        expected_columns=[("id", "int"), ("nome", "str"), ("valor", "int")],
        notes="A 2-column metadata block + blank line precede the real 3-column table.",
    ),
    Fixture(
        name="f04_duplicate_headers.csv",
        feature="F2",
        build=lambda p: _write_csv(p, "nome,valor,valor\nAna,1,2\nBeto,3,4\n"),
        expected_data=[
            {"nome": "Ana", "valor": 1, "valor_2": 2},
            {"nome": "Beto", "valor": 3, "valor_2": 4},
        ],
        expected_columns=[("nome", "str"), ("valor", "int"), ("valor_2", "int")],
        notes="Duplicate 'valor' header → valor, valor_2.",
    ),
    Fixture(
        name="f05_empty_header.csv",
        feature="F2",
        build=lambda p: _write_csv(p, "id,,desc\n1,x,y\n2,z,w\n"),
        expected_data=[
            {"id": 1, "column_2": "x", "desc": "y"},
            {"id": 2, "column_2": "z", "desc": "w"},
        ],
        expected_columns=[("id", "int"), ("column_2", "str"), ("desc", "str")],
        notes="Empty header cell → synthesised column_2.",
    ),
    Fixture(
        name="f06_multiline_header.xlsx",
        feature="F2",
        build=lambda p: _write_xlsx(
            p,
            [
                ["Produto", "Vendas", None, "Estoque"],
                [None, "2023", "2024", None],
                ["A", 100, 150, 5],
                ["B", 200, 250, 8],
            ],
            merges=("B1:C1",),
        ),
        expected_data=[
            {"produto": "A", "vendas_2023": 100, "vendas_2024": 150, "estoque": 5},
            {"produto": "B", "vendas_2023": 200, "vendas_2024": 250, "estoque": 8},
        ],
        expected_columns=[
            ("produto", "str"),
            ("vendas_2023", "int"),
            ("vendas_2024", "int"),
            ("estoque", "int"),
        ],
        notes="Merged group header 'Vendas' over '2023'/'2024' → vendas_2023, vendas_2024.",
    ),
    Fixture(
        name="f07_merged_fill.xlsx",
        feature="F3",
        build=lambda p: _write_xlsx(
            p,
            [
                ["regiao", "produto", "valor"],
                ["Sul", "A", 10],
                [None, "B", 20],
                ["Norte", "C", 30],
            ],
            merges=("A2:A3",),
        ),
        expected_data=[
            {"regiao": "Sul", "produto": "A", "valor": 10},
            {"regiao": "Sul", "produto": "B", "valor": 20},
            {"regiao": "Norte", "produto": "C", "valor": 30},
        ],
        expected_columns=[("regiao", "str"), ("produto", "str"), ("valor", "int")],
        notes="Vertically merged 'Sul' propagated down (fill mode).",
    ),
    Fixture(
        name="f08_merged_first_only.xlsx",
        feature="F3",
        build=lambda p: _write_xlsx(
            p,
            [
                ["regiao", "produto", "valor"],
                ["Sul", "A", 10],
                [None, "B", 20],
                ["Norte", "C", 30],
            ],
            merges=("A2:A3",),
        ),
        config=Config(merged_cells="first-only"),
        expected_data=[
            {"regiao": "Sul", "produto": "A", "valor": 10},
            {"regiao": None, "produto": "B", "valor": 20},
            {"regiao": "Norte", "produto": "C", "valor": 30},
        ],
        expected_columns=[("regiao", "str"), ("produto", "str"), ("valor", "int")],
        notes="Same file, first-only mode → merged value stays top-left.",
    ),
    Fixture(
        name="f09_serial_dates.xlsx",
        feature="F4",
        build=lambda p: _write_xlsx(
            p,
            [
                ["item", "data_venda", "preco"],
                ["X", 45123, 9.9],
                ["Y", 45124, 8.5],
            ],
        ),
        expected_data=[
            {"item": "X", "data_venda": _D(2023, 7, 16), "preco": 9.9},
            {"item": "Y", "data_venda": _D(2023, 7, 17), "preco": 8.5},
        ],
        expected_columns=[("item", "str"), ("data_venda", "date"), ("preco", "float")],
        notes="Bare serials in a 'data_*' column → real dates (1900 epoch).",
    ),
    Fixture(
        name="f10_numbers_ptbr.csv",
        feature="F5",
        build=lambda p: _write_csv(p, "item;valor\nA;1.234,56\nB;2.000,00\nC;42\n"),
        expected_data=[
            {"item": "A", "valor": 1234.56},
            {"item": "B", "valor": 2000.0},
            {"item": "C", "valor": 42.0},
        ],
        expected_columns=[("item", "str"), ("valor", "float")],
        notes="pt-BR numbers with ';' delimiter; column-level locale inference.",
    ),
    Fixture(
        name="f11_numbers_enus.csv",
        feature="F5",
        build=lambda p: _write_csv(p, 'item,price\nA,"1,234.56"\nB,"2,000.00"\n'),
        expected_data=[
            {"item": "A", "price": 1234.56},
            {"item": "B", "price": 2000.0},
        ],
        expected_columns=[("item", "str"), ("price", "float")],
        notes="en-US numbers, comma thousands quoted to survive the comma delimiter.",
    ),
    Fixture(
        name="f12_mixed_numeric_nulls.csv",
        feature="F6",
        build=lambda p: _write_csv(p, "produto,qtd\nA,10\nB,20\nC,N/A\nD,30\n"),
        expected_data=[
            {"produto": "A", "qtd": 10},
            {"produto": "B", "qtd": 20},
            {"produto": "C", "qtd": None},
            {"produto": "D", "qtd": 30},
        ],
        expected_columns=[("produto", "str"), ("qtd", "int")],
        notes="Mostly-numeric column with an N/A → numeric with a reported null.",
    ),
    Fixture(
        name="f13_disguised_nulls.csv",
        feature="F7",
        build=lambda p: _write_csv(p, "id,valor,nota\n1,10,-\n2,N/A,#REF!\n3,#DIV/0!,ok\n"),
        expected_data=[
            {"id": 1, "valor": 10, "nota": None},
            {"id": 2, "valor": None, "nota": None},
            {"id": 3, "valor": None, "nota": "ok"},
        ],
        expected_columns=[("id", "int"), ("valor", "int"), ("nota", "str")],
        notes="-, N/A, #REF!, #DIV/0! all normalised to null.",
    ),
    Fixture(
        name="f14_trailing_totals.xlsx",
        feature="F8",
        build=lambda p: _write_xlsx(
            p,
            [
                ["produto", "valor"],
                ["A", 10],
                ["B", 20],
                ["TOTAL", 30],
                [None, None],
                ["Gerado em 01/02/2024", None],
            ],
        ),
        expected_data=[{"produto": "A", "valor": 10}, {"produto": "B", "valor": 20}],
        expected_columns=[("produto", "str"), ("valor", "int")],
        notes="TOTAL row, blank row and a 'Gerado em' footnote trimmed.",
    ),
    Fixture(
        name="f15_latin1_semicolon.csv",
        feature="F10",
        build=lambda p: _write_csv(
            p,
            "Município;População\nSão Paulo;12000000\nRio;6000000\n",
            encoding="latin-1",
        ),
        expected_data=[
            {"municipio": "São Paulo", "populacao": 12000000},
            {"municipio": "Rio", "populacao": 6000000},
        ],
        expected_columns=[("municipio", "str"), ("populacao", "int")],
        notes="Latin-1 encoded, ';' delimiter, accented headers.",
    ),
    Fixture(
        name="f16_basic.tsv",
        feature="F10",
        build=lambda p: _write_csv(p, "a\tb\tc\n1\tx\t3.5\n2\ty\t4.5\n"),
        expected_data=[
            {"a": 1, "b": "x", "c": 3.5},
            {"a": 2, "b": "y", "c": 4.5},
        ],
        expected_columns=[("a", "int"), ("b", "str"), ("c", "float")],
        notes="Tab-separated values.",
    ),
    Fixture(
        name="f17_no_header.csv",
        feature="F2",
        build=lambda p: _write_csv(p, "10,20,30\n40,50,60\n"),
        config=Config(header=None),
        expected_data=[
            {"column_1": 10, "column_2": 20, "column_3": 30},
            {"column_1": 40, "column_2": 50, "column_3": 60},
        ],
        expected_columns=[("column_1", "int"), ("column_2", "int"), ("column_3", "int")],
        notes="header=None → synthesised column names, first row is data.",
    ),
]
