"""Reader edge cases: encoding, delimiter, formats, sheets, bad input."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from _fixtures import _write_csv, _write_xlsx
from messy_table import Config, MessyTableError, UnsupportedFormatError, clean


def test_utf8_with_bom(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "bom.csv", "nome,v\nÁgua,1\n", encoding="utf-8-sig")
    result = clean(path)
    assert result.data == [{"nome": "Água", "v": 1}]
    assert result.source.encoding == "utf-8-sig"


def test_cp1252_fallback(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "win.csv", "nome,v\nAção,1\n", encoding="cp1252")
    result = clean(path)
    assert result.data[0]["nome"] == "Ação"
    assert result.source.encoding in {"cp1252", "latin-1"}


def test_pipe_delimiter(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "p.csv", "a|b|c\n1|x|2\n3|y|4\n")
    result = clean(path)
    assert result.source.delimiter == "pipe"
    assert result.data == [{"a": 1, "b": "x", "c": 2}, {"a": 3, "b": "y", "c": 4}]


def test_tsv_delimiter(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "t.tsv", "a\tb\n1\tx\n")
    result = clean(path)
    assert result.source.delimiter == "tab"


def test_bytes_xlsx_detected_by_magic() -> None:
    buf = io.BytesIO()
    wb = Workbook()
    wb.active.append(["x", "y"])
    wb.active.append([1, 2])
    wb.save(buf)
    result = clean(buf.getvalue())  # no extension; sniffed via PK magic
    assert result.data == [{"x": 1, "y": 2}]


def test_filelike_text_input() -> None:
    result = clean(io.StringIO("a,b\n1,2\n"))
    assert result.data == [{"a": 1, "b": 2}]


def test_xlsx_sheet_by_name(tmp_path: Path) -> None:
    wb = Workbook()
    wb.active.title = "first"
    wb.active.append(["x"])
    wb.active.append([1])
    second = wb.create_sheet("second")
    second.append(["y"])
    second.append([9])
    path = tmp_path / "multi.xlsx"
    wb.save(path)
    result = clean(path, config=Config(sheet="second"))
    assert result.data == [{"y": 9}]


def test_xlsx_sheet_not_found(tmp_path: Path) -> None:
    path = _write_xlsx(tmp_path / "one.xlsx", [["x"], [1]])
    with pytest.raises(MessyTableError, match="not found"):
        clean(path, config=Config(sheet="nope"))


def test_unsupported_xls(tmp_path: Path) -> None:
    path = tmp_path / "data.xls"
    path.write_bytes(b"whatever")
    with pytest.raises(UnsupportedFormatError, match="legacy"):
        clean(path)


def test_unsupported_ods(tmp_path: Path) -> None:
    path = tmp_path / "data.ods"
    path.write_bytes(b"whatever")
    with pytest.raises(UnsupportedFormatError, match="ods"):
        clean(path)


def test_file_not_found() -> None:
    with pytest.raises(UnsupportedFormatError, match="not found"):
        clean("/no/such/file.csv")


def test_unsupported_source_type() -> None:
    with pytest.raises(UnsupportedFormatError, match="unsupported source"):
        clean(12345)  # type: ignore[arg-type]


def test_empty_input_returns_empty_result(tmp_path: Path) -> None:
    path = _write_csv(tmp_path / "empty.csv", "\n")  # .csv ext bypasses content sniff
    result = clean(path)
    assert result.data == []
    assert any("no data" in w.message for w in result.warnings)


def test_undeterminable_format_raises() -> None:
    with pytest.raises(UnsupportedFormatError, match="determine format"):
        clean(b"   ")  # whitespace-only, no extension
